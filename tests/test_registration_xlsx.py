"""Offline checks for optional registration Team columns."""

import contextlib
import io
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from libs.racedb import RaceDB


class RegistrationXlsxTests(unittest.TestCase):
    def test_optional_team_column(self):
        racedb = RaceDB.__new__(RaceDB)
        original_headers = RaceDB.RegistrationHeaders.copy()
        rows = [{'First Name': 'Test', 'Last Name': 'Rider', 'Team': 'Test Club'}]
        with tempfile.TemporaryDirectory() as directory:
            for in_memory in (False, True):
                for teams in (True, False, True):
                    with self.subTest(in_memory=in_memory, teams=teams):
                        if in_memory:
                            source = racedb.create_registrations_xlsx_in_memory(rows, teams=teams)
                        else:
                            source = Path(directory) / 'registrations.xlsx'
                            with contextlib.redirect_stdout(io.StringIO()):
                                racedb.create_registrations_xlsx_file(rows, source, teams=teams)
                        workbook = load_workbook(source)
                        try:
                            sheet = workbook.active
                            headers = [cell.value for cell in sheet[1]]
                            self.assertEqual(headers, list(original_headers) + (['Team'] if teams else []))
                            if teams:
                                self.assertEqual(sheet.cell(2, headers.index('Team') + 1).value, 'Test Club')
                            self.assertEqual(RaceDB.RegistrationHeaders, original_headers)
                        finally:
                            workbook.close()
                            if in_memory:
                                source.close()


if __name__ == '__main__':
    unittest.main()
