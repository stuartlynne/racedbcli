import io
from pathlib import Path
import sys
import unittest
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from crossmgr_fix_teams import match_holder, patch_workbook


class CrossMgrTeamTests(unittest.TestCase):
    def test_identity_requires_unique_name_or_matching_license(self):
        row = {'last_name': 'Smith', 'first_name': 'Alex', 'license': 'TEMP'}
        holder = {'id': 1, 'last_name': 'Smith', 'first_name': 'Alex',
                  'license_code': 'BC123', 'team_source': 'discipline default'}
        self.assertEqual(match_holder(row, [holder])[0], holder)
        other = dict(holder, id=2, license_code='BC456')
        self.assertIsNone(match_holder(row, [holder, other])[0])
        row['license'] = 'BC456'
        self.assertEqual(match_holder(row, [holder, other])[0], other)
        row['license'] = 'BC789'
        self.assertIsNone(match_holder(row, [holder])[0])

    def test_only_team_cells_and_registration_zip_part_change(self):
        w = Workbook()
        s = w.active
        s.title = 'Registration'
        s.append(['LastName', 'FirstName', 'Team', 'License'])
        s.append(['Smith', 'Alex', 'Old Team', 'BC123'])
        s.append(['Jones', 'Kim', None, 'BC456'])
        s.freeze_panes = 'C2'
        w.create_sheet('--CrossMgr-Categories').append(['Keep categories', '=1+2'])
        w.create_sheet('--CrossMgr-Properties').append(['Keep properties', True])
        raw = io.BytesIO()
        w.save(raw)
        updated, part = patch_workbook(raw.getvalue(), {'C2': 'A & B', 'C3': '=literal'})
        with ZipFile(raw) as old, ZipFile(io.BytesIO(updated)) as new:
            for name in old.namelist():
                if name != part:
                    self.assertEqual(old.read(name), new.read(name), name)
        result = load_workbook(io.BytesIO(updated))
        self.assertEqual(result.sheetnames, w.sheetnames)
        self.assertEqual(result['Registration']['C2'].value, 'A & B')
        self.assertEqual(result['Registration']['C3'].value, '=literal')
        self.assertEqual(result['Registration']['C3'].data_type, 's')
        self.assertEqual(result['Registration']['D3'].value, 'BC456')
        self.assertEqual(result['Registration'].freeze_panes, 'C2')


if __name__ == '__main__':
    unittest.main()
