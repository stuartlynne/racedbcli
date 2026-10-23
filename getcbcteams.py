#!/usr/bin/env python3
"""Export Cycling BC teams to XLSX (curl, beautifulsoup4, openpyxl required)."""

import argparse
from collections import Counter
from datetime import datetime
from pathlib import Path
import re
import subprocess
import sys
from urllib.parse import urljoin

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


SOURCE_URL = "https://cyclingbc.net/clubs/affiliates/"
TABLE_NAMES = (
    "Provincial Clubs", "Provincial Trade Teams", "National Teams", "UCI Teams",
)


def clean_text(element):
    return " ".join(element.get_text(" ", strip=True).split())


def fetch_page(url):
    # This site currently accepts curl but rejects Python requests with HTTP 403.
    result = subprocess.run(
        ["curl", "--fail", "--silent", "--show-error", "--location",
         "--connect-timeout", "15", "--max-time", "60", url],
        capture_output=True, check=True,
    )
    return result.stdout


def parse_page(html, source_url=SOURCE_URL):
    soup = BeautifulSoup(html, "html.parser")
    match = re.search(
        r"last updated on\s+([A-Za-z]+\s+\d{1,2},\s+\d{4})",
        clean_text(soup), re.IGNORECASE,
    )
    if not match:
        raise ValueError("Page has no recognizable 'last updated on' date")
    for date_format in ("%b %d, %Y", "%B %d, %Y"):
        try:
            updated = datetime.strptime(match.group(1), date_format).date()
            break
        except ValueError:
            continue
    else:
        raise ValueError(f"Unrecognized update date: {match.group(1)}")

    headers = ["Club / Team", "Table"]
    records = []
    for table_name in TABLE_NAMES:
        tab = next((a for a in soup.select('a[href^="#"]')
                    if clean_text(a) == table_name), None)
        panel = soup.find(id=tab["href"][1:]) if tab else None
        table = panel.find("table") if panel else None
        if table is None:
            raise ValueError(f"Missing table: {table_name}")
        columns = [clean_text(cell) for cell in table.select("thead th")]
        if not columns or not {"Club", "Team"}.intersection(columns):
            raise ValueError(f"Unrecognized headers in {table_name}: {columns}")
        for column in columns:
            key = "Club / Team" if column in ("Club", "Team") else column
            if key not in headers:
                headers.append(key)
        region = ""
        for row in table.select("tbody tr"):
            cells = row.find_all(["td", "th"], recursive=False)
            if len(cells) != len(columns):
                raise ValueError(f"Unexpected row width in {table_name}: {clean_text(row)}")
            record = {"Table": table_name}
            for column, cell in zip(columns, cells):
                key = "Club / Team" if column in ("Club", "Team") else column
                record[key] = clean_text(cell)
                urls = []
                for anchor in cell.select("a[href]"):
                    href = anchor["href"].strip()
                    if not href:
                        continue
                    url = "https://" + href if href.startswith("www.") else urljoin(source_url, href)
                    if url not in urls:
                        urls.append(url)
                if urls:
                    url_key = key + " URL"
                    record[url_key] = "\n".join(urls)
                    if url_key not in headers:
                        headers.append(url_key)
            region = record.get("Region") or region
            if "Region" in columns:
                record["Region"] = region
            # Region-only rows label a group; they are not separate clubs.
            if not record.get("Club / Team"):
                if any(value for key, value in record.items() if key not in ("Region", "Table")):
                    raise ValueError(f"Unnamed entry in {table_name}: {record}")
                continue
            record["Source URL"] = source_url
            records.append(record)
    headers.append("Source URL")
    if not records:
        raise ValueError("No teams found")
    records.sort(key=lambda row: (row["Club / Team"].casefold(), row["Table"]))
    return updated, headers, records


def write_workbook(path, updated, headers, records):
    workbook = Workbook()
    workbook.properties.title = "Cycling BC affiliated clubs and teams"
    workbook.properties.description = f"Source last updated: {updated.isoformat()}"
    sheet = workbook.active
    sheet.title = "Clubs and Teams"
    sheet.append(headers)
    for record in records:
        sheet.append([record.get(header, "") for header in headers])
    for row in sheet:
        for cell in row:
            # Preserve scraped text literally, including anything starting with '='.
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row > 1 and headers[cell.column - 1].endswith("URL"):
                value = cell.value or ""
                if value.startswith(("https://", "http://")) and "\n" not in value:
                    cell.hyperlink = value
                    cell.style = "Hyperlink"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17465B")
    sheet.row_dimensions[1].height = 30
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 60
    for index, header in enumerate(headers, 1):
        width = 45 if header.endswith("URL") else {
            "Club / Team": 44, "Table": 27, "Region": 29, "City": 22,
            "Disciplines": 40, "Website": 34, "Portal": 12,
        }.get(header, 25)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    table = Table(displayName="CBCTeams", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True,
    )
    sheet.add_table(table)
    sheet.auto_filter.ref = sheet.dimensions
    sheet.auto_filter.add_sort_condition(f"A2:A{sheet.max_row}")
    workbook.save(path)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", default=SOURCE_URL, help="Affiliate page URL")
    parser.add_argument("--html", type=Path, help="Read a saved page instead of downloading")
    parser.add_argument("--output", type=Path, help="Default: cbcteams-YYYY-MM-DD.xlsx using page update date")
    args = parser.parse_args(argv)
    try:
        html = args.html.read_bytes() if args.html else fetch_page(args.url)
        updated, headers, records = parse_page(html, args.url)
        output = args.output or Path(f"cbcteams-{updated.isoformat()}.xlsx")
        write_workbook(output, updated, headers, records)
    except subprocess.CalledProcessError as exc:
        print(f"Download failed: {exc.stderr.decode(errors='replace').strip()}", file=sys.stderr)
        return 1
    except (OSError, ValueError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    for name, count in Counter(record["Table"] for record in records).items():
        print(f"{name}: {count}", file=sys.stderr)
    print(f"Wrote {len(records)} entries to {output} (source updated {updated})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
