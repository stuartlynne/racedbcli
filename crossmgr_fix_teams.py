#!/usr/bin/env python3
"""Update CrossMgr Registration teams from RaceDB, preserving other XLSX parts.

Requires mcp, openpyxl, and lxml. Reads RaceDB through MCP; never writes to RaceDB.
"""

import argparse
import asyncio
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import posixpath
import re
import sys
import tempfile
import unicodedata
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter

from libs.mcpconfig import load_mcp_host

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def normalize(value):
    text = unicodedata.normalize('NFKD', str(value or '').casefold())
    return ' '.join(''.join(c for c in text if not unicodedata.combining(c)).split())


def quote(value):
    return "'" + str(value).replace("'", "''") + "'"


def usable_license(value):
    return normalize(value) not in ('', 'temp', 'none', 'n/a')


def inspect_file(path):
    raw = path.read_bytes()
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    try:
        sheet = workbook['Registration']
        headers = {str(c.value): c.column for c in next(sheet.iter_rows())}
        for name in ('LastName', 'FirstName', 'Team'):
            if name not in headers:
                raise ValueError(f'{path}: missing {name} header')
        rows = []
        for row in sheet.iter_rows(min_row=2):
            get = lambda key: row[headers[key] - 1].value if key in headers else None
            if not get('LastName') and not get('FirstName'):
                continue
            rows.append({'row': row[0].row, 'last_name': get('LastName'),
                         'first_name': get('FirstName'), 'license': get('License'),
                         'old_team': get('Team'),
                         'cell': f"{get_column_letter(headers['Team'])}{row[0].row}"})
        return {'path': path, 'raw': raw, 'rows': rows}
    finally:
        workbook.close()


def team_query(last_names, licenses, discipline_id):
    filters = ['lower(l.last_name) IN (' + ','.join(map(quote, last_names)) + ')']
    if licenses:
        filters.append('l.license_code IN (' + ','.join(map(quote, licenses)) + ')')
    # Mirror LicenseHolder.get_team_for_discipline, including its fallback.
    return f"""SELECT l.id,l.last_name,l.first_name,l.license_code,t.name AS team,
        CASE WHEN p.id IS NOT NULL THEN 'registration'
             WHEN h.id IS NOT NULL THEN 'discipline default'
             WHEN a.id IS NOT NULL THEN 'other discipline registration'
             ELSE NULL END AS team_source
        FROM core_licenseholder l
        LEFT JOIN LATERAL (
            SELECT id,team_id,effective_date FROM core_teamhint
            WHERE license_holder_id=l.id AND discipline_id={discipline_id}
            ORDER BY effective_date DESC,id DESC LIMIT 1
        ) h ON TRUE
        LEFT JOIN LATERAL (
            SELECT p.id,p.team_id FROM core_participant p
            JOIN core_competition c ON c.id=p.competition_id
            WHERE p.license_holder_id=l.id AND c.discipline_id={discipline_id}
              AND (h.id IS NULL OR c.start_date>h.effective_date)
            ORDER BY c.start_date DESC,p.id DESC LIMIT 1
        ) p ON TRUE
        LEFT JOIN LATERAL (
            SELECT p.id,p.team_id FROM core_participant p
            JOIN core_competition c ON c.id=p.competition_id
            WHERE p.license_holder_id=l.id AND h.id IS NULL
            ORDER BY c.start_date DESC,p.id DESC LIMIT 1
        ) a ON h.id IS NULL AND p.id IS NULL
        LEFT JOIN core_team t ON t.id=CASE WHEN p.id IS NOT NULL THEN p.team_id
            WHEN h.id IS NOT NULL THEN h.team_id ELSE a.team_id END
        WHERE {' OR '.join(filters)} ORDER BY l.id"""


async def fetch_teams(config, rows, discipline, token):
    from mcp import ClientSession
    from mcp.client.streamable_http import streamablehttp_client

    async with streamablehttp_client(config['mcp_url'],
            headers={'Authorization': f'Bearer {token}'} if token else {}) as (read, write, _):
        async with ClientSession(read, write) as session:
            await session.initialize()
            listed = []
            cursor = None
            while True:
                page = await session.list_tools(cursor=cursor)
                listed.extend(t.name for t in page.tools)
                cursor = page.nextCursor
                if not cursor:
                    break
            matches = [n for n in listed if re.sub('[^a-z0-9]', '', n.lower()).endswith('sqlquery')]
            tool = 'sql.query' if 'sql.query' in listed else matches[0] if len(matches) == 1 else None
            if not tool:
                raise RuntimeError('Cannot identify a unique sql.query tool')

            async def query(sql):
                result = await session.call_tool(tool, arguments={'sql': sql, 'limit': 1000})
                if result.isError:
                    raise RuntimeError(str(result.content))
                data = result.structuredContent
                if not isinstance(data, dict):
                    data = next((json.loads(c.text) for c in result.content if c.type == 'text'), {})
                if 'rows' not in data:
                    raise RuntimeError('SQL response missing rows')
                return data

            data = await query('SELECT id,name FROM core_discipline WHERE lower(name)=' + quote(discipline.lower()))
            if len(data['rows']) != 1:
                raise RuntimeError(f'Unknown/ambiguous discipline: {discipline}')
            did = int(data['rows'][0]['id'])
            names = sorted({str(r['last_name']).strip().lower() for r in rows if r['last_name']})
            licenses = sorted({str(r['license']).strip() for r in rows if usable_license(r['license'])})
            sql = team_query(names, licenses, did)
            found = []
            offset = 0
            while True:
                data = await query(sql + f' LIMIT 1000 OFFSET {offset}')
                found.extend(data['rows'])
                if len(data['rows']) < 1000:
                    break
                offset += 1000
            return found


def match_holder(row, holders):
    matches = [h for h in holders if normalize(h['last_name']) == normalize(row['last_name'])
               and normalize(h['first_name']) == normalize(row['first_name'])]
    if usable_license(row['license']):
        matches = [h for h in matches if normalize(h['license_code']) == normalize(row['license'])]
    if len(matches) != 1:
        return None, 'ambiguous identity' if len(matches) > 1 else 'no matching identity'
    if not matches[0]['team_source']:
        return None, 'no team history; retained existing value'
    return matches[0], ''


def patch_workbook(raw, changes):
    with ZipFile(io.BytesIO(raw)) as source:
        book = etree.fromstring(source.read('xl/workbook.xml'))
        sheet = next(s for s in book.find('s:sheets', NS) if s.get('name') == 'Registration')
        rid = sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        rels = etree.fromstring(source.read('xl/_rels/workbook.xml.rels'))
        target = next(r.get('Target') for r in rels if r.get('Id') == rid)
        part = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        root = etree.fromstring(source.read(part))
        sheet_data = root.find('s:sheetData', NS)
        for address, text in changes.items():
            row_number = re.search(r'\d+$', address).group()
            row = next(r for r in sheet_data if r.get('r') == row_number)
            cell = next((c for c in row if c.get('r') == address), None)
            if cell is None:
                cell = etree.Element('{%s}c' % NS['s'], r=address)
                column = column_index_from_string(re.sub(r'\d+', '', address))
                index = next((i for i, c in enumerate(row)
                    if column_index_from_string(re.sub(r'\d+', '', c.get('r'))) > column), len(row))
                row.insert(index, cell)
            for child in list(cell):
                cell.remove(child)
            cell.set('t', 'inlineStr')
            text_element = etree.SubElement(etree.SubElement(cell, '{%s}is' % NS['s']), '{%s}t' % NS['s'])
            text_element.text = text
            text_element.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        destination = io.BytesIO()
        with ZipFile(destination, 'w') as out:
            out.comment = source.comment
            for info in source.infolist():
                out.writestr(info, etree.tostring(root) if info.filename == part else source.read(info.filename))
        return destination.getvalue(), part


def save_changes(item, changes):
    path, raw = item['path'], item['raw']
    updated, part = patch_workbook(raw, changes)
    with ZipFile(io.BytesIO(raw)) as before, ZipFile(io.BytesIO(updated)) as after:
        assert before.namelist() == after.namelist()
        assert all(before.read(n) == after.read(n) for n in before.namelist() if n != part)
    workbook = load_workbook(io.BytesIO(updated), read_only=True)
    try:
        assert all(workbook['Registration'][c].value == v for c, v in changes.items())
    finally:
        workbook.close()
    temp = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, suffix='.xlsx', delete=False) as f:
            temp = Path(f.name)
            f.write(updated)
        temp.chmod(path.stat().st_mode & 0o777)
        if hashlib.sha256(path.read_bytes()).digest() != hashlib.sha256(raw).digest():
            raise RuntimeError(f'{path} changed during processing; refusing to overwrite')
        os.replace(temp, path)
    finally:
        if temp and temp.exists():
            temp.unlink()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('files', nargs='+', type=Path)
    parser.add_argument('--mcp-host', required=True, help='Host key in ~/.mcp.json')
    parser.add_argument('--mcp-config')
    parser.add_argument('--discipline', default='Cyclocross')
    parser.add_argument('--dry-run', action='store_true', help='Report proposed changes without saving files')
    parser.add_argument('--report', type=Path, default=Path('crossmgr-team-updates.csv'))
    args = parser.parse_args()
    config = load_mcp_host(args.mcp_host, args.mcp_config)
    if not config['mcp_url']:
        parser.error('Selected host has no mcp_url')
    items = [inspect_file(p.resolve()) for p in dict.fromkeys(args.files)]
    rows = [r for item in items for r in item['rows']]
    if not rows:
        parser.error('No registrants found')
    holders = asyncio.run(fetch_teams(config, rows, args.discipline, os.getenv('CONTEXTFORGE_TOKEN') or os.getenv('RACEDB_MCP_TOKEN')))
    report = []
    for item in items:
        changes = {}
        for row in item['rows']:
            holder, reason = match_holder(row, holders)
            new_team = (holder['team'] or 'Independent') if holder else row['old_team']
            status = 'skipped' if not holder else 'unchanged' if new_team == row['old_team'] else 'would update' if args.dry_run else 'updated'
            if holder and new_team != row['old_team']:
                changes[row['cell']] = new_team
            report.append({'file': str(item['path']), **row, 'new_team': new_team,
                           'holder_id': holder['id'] if holder else '', 'status': status,
                           'reason': reason or holder['team_source']})
        if changes and not args.dry_run:
            save_changes(item, changes)
        print(f"{item['path'].name}: {len(changes)} {'proposed' if args.dry_run else 'updated'} / {len(item['rows'])} rows")
    if report:
        with args.report.open('w', newline='', encoding='utf-8-sig') as f:
            writer = csv.DictWriter(f, fieldnames=list(report[0]))
            writer.writeheader()
            writer.writerows(report)
    print(f"Skipped {sum(r['status']=='skipped' for r in report)} rows; report: {args.report}")


if __name__ == '__main__':
    main()
