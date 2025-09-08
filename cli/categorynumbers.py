#!/usr/bin/env python3

import sys
import os
import argparse
from typing import Dict, List, Tuple, Optional
import re

# repo root for libs import
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from libs.racedbsql import RaceDBSQL  # type: ignore


def find_competition(db: RaceDBSQL, name: Optional[str], date: Optional[str]):
    comp, _ = db.find_competition_categories(name=name, date=date)
    if not comp:
        raise SystemExit("Competition not found (use --name or --date)")
    return comp


def fetch_categorynumbers(db: RaceDBSQL, competition_id: int) -> List[Dict]:
    q = (
        "SELECT cn.id AS cn_id, cn.range_str, c.code AS category_code, c.gender AS category_gender "
        "FROM core_categorynumbers cn "
        "JOIN core_categorynumbers_categories cnc ON cnc.categorynumbers_id = cn.id "
        "JOIN core_category c ON c.id = cnc.category_id "
        "WHERE cn.competition_id = %s "
        "ORDER BY c.code, cn.id"
    )
    db.cur_execute("Fetch categorynumbers", q, (competition_id,), debug=False)
    return db.cur.fetchall() or []


def get_code_gender_to_id(db: RaceDBSQL, format_id: int) -> Dict[Tuple[str, int], int]:
    q = "SELECT id, code, gender FROM core_category WHERE format_id = %s;"
    db.cur_execute("Fetch categories for format", q, (format_id,), debug=False)
    rows = db.cur.fetchall() or []
    mapping: Dict[Tuple[str, int], int] = {}
    for r in rows:
        cid = r.get("id") if isinstance(r, dict) else r[0]
        code = r.get("code") if isinstance(r, dict) else r[1]
        gender = r.get("gender") if isinstance(r, dict) else r[2]
        mapping[(str(code), int(gender) if gender is not None else -1)] = int(cid)
    return mapping


def get_id_to_code_gender(db: RaceDBSQL, format_id: int) -> Dict[int, Tuple[str, int]]:
    q = "SELECT id, code, gender FROM core_category WHERE format_id = %s;"
    db.cur_execute("Fetch id->(code,gender) map", q, (format_id,), debug=False)
    rows = db.cur.fetchall() or []
    mapping: Dict[int, Tuple[str, int]] = {}
    for r in rows:
        cid = r.get("id") if isinstance(r, dict) else r[0]
        code = r.get("code") if isinstance(r, dict) else r[1]
        gender = r.get("gender") if isinstance(r, dict) else r[2]
        mapping[int(cid)] = (str(code), int(gender) if gender is not None else 2)
    return mapping


def fetch_categories_for_format(db: RaceDBSQL, format_id: int) -> List[Dict]:
    q = "SELECT id, code, gender, sequence FROM core_category WHERE format_id = %s ORDER BY sequence, code;"
    db.cur_execute("Fetch all categories in format", q, (format_id,), debug=False)
    return db.cur.fetchall() or []


def sanitize_filename(name: str) -> str:
    # Replace spaces with underscores and remove invalid characters
    s = name.replace(" ", "_")
    # Allow letters, digits, underscore, hyphen, dot
    s = re.sub(r"[^A-Za-z0-9._-]", "", s)
    # Collapse multiple underscores
    s = re.sub(r"_+", "_", s)
    return s or "competition"


def normalize_ranges(range_list: List[str]) -> str:
    parts = [p.strip() for p in range_list if str(p).strip()]
    # keep order but deduplicate
    seen = set()
    norm = []
    for p in parts:
        if p not in seen:
            seen.add(p)
            norm.append(p)
    return ",".join(norm)


def fetch_race_category_map(db: RaceDBSQL, competition_id: int, debug: bool = False) -> Dict[Tuple[str, int], Tuple[int, int]]:
    """Return mapping from (code, gender) to (race_no, start_offset).

    race_no increments by event (mass start) ordered by event start time.
    start_offset is the wave start offset within an event (None treated as 0).
    """
    q = (
        "SELECT e.id AS event_id, e.date_time AS event_start_time, w.start_offset, c.code, c.gender "
        "FROM core_eventmassstart e "
        "LEFT JOIN core_wave w ON w.event_id = e.id "
        "JOIN core_wave_categories wc ON wc.wave_id = w.id "
        "JOIN core_category c ON wc.category_id = c.id "
        "WHERE e.competition_id = %s "
        "ORDER BY e.date_time, w.start_offset, c.code;"
    )
    db.cur_execute("Fetch race-sorted categories", q, (competition_id,), debug=False)
    rows = db.cur.fetchall() or []

    # Determine race numbers by unique event_id in event_start_time order
    event_order: List[int] = []
    race_no_by_event: Dict[int, int] = {}
    for r in rows:
        event_id = r.get("event_id") if isinstance(r, dict) else r[0]
        if event_id not in race_no_by_event:
            event_order.append(event_id)
            race_no_by_event[event_id] = len(event_order)

    mapping: Dict[Tuple[str, int], Tuple[int, int]] = {}
    last_event = None
    last_offset = None
    for r in rows:
        event_id = r.get("event_id") if isinstance(r, dict) else r[0]
        event_time = r.get("event_start_time") if isinstance(r, dict) else r[1]
        start_offset = r.get("start_offset") if isinstance(r, dict) else r[2]
        code = r.get("code") if isinstance(r, dict) else r[3]
        gender = r.get("gender") if isinstance(r, dict) else r[4]
        key = (str(code), int(gender) if gender is not None else -1)
        # First seen wins; map to race number and start offset
        if key not in mapping:
            mapping[key] = (race_no_by_event.get(event_id, 10**9), int(start_offset) if start_offset is not None else 0)

        if debug:
            # Print grouped debug: races (events), then starts (offset), then categories
            race_no = race_no_by_event.get(event_id, 10**9)
            if last_event != event_id:
                print(f"Race {race_no}: event_id={event_id}, start_time={event_time}", file=sys.stderr)
                last_event = event_id
                last_offset = None
            if last_offset != start_offset:
                print(f"  Start offset: {start_offset}", file=sys.stderr)
                last_offset = start_offset
            gstr = 'Men' if (gender or 2) == 0 else ('Women' if (gender or 2) == 1 else 'Open')
            print(f"    Category: {code} ({gstr})", file=sys.stderr)
    return mapping


def download_xlsx(db: RaceDBSQL, comp: Dict, output_path: str, sort_by_race: bool = False, debug: bool = False) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        from openpyxl.styles import Alignment
    except Exception as e:
        raise SystemExit(f"openpyxl not available: {e}")

    rows = fetch_categorynumbers(db, comp["id"])
    # Collect rows first to determine header range count
    present_codes: List[str] = []
    data_rows: List[List[str]] = []

    def gender_to_str(gval: Optional[int]) -> str:
        try:
            gv = int(gval) if gval is not None else 2
        except Exception:
            gv = 2
        return 'Men' if gv == 0 else ('Women' if gv == 1 else 'Open')

    # Build race mapping for display
    race_map: Dict[Tuple[str, int], Tuple[int, int]] = {}
    try:
        race_map = fetch_race_category_map(db, comp["id"], debug=debug)  # (code,g) -> (race_no, start_offset)
    except Exception:
        race_map = {}

    for r in rows:
        code = r.get("category_code") if isinstance(r, dict) else r[2]
        gval = r.get("category_gender") if isinstance(r, dict) else r[3]
        display_code = f"{code} ({gender_to_str(gval)})"
        range_str = r.get("range_str") if isinstance(r, dict) else r[1]
        ranges = [p.strip() for p in str(range_str or "").split(",") if p.strip()]
        race_key = (str(code), int(gval) if gval is not None else -1)
        race_no = race_map.get(race_key, ("", 0))[0]
        data_rows.append([race_no, display_code] + ranges)
        present_codes.append((str(code), int(gval) if gval is not None else -1))

    # Ensure a line for each category in the format, even if no ranges
    all_cats = fetch_categories_for_format(db, comp["category_format_id"])
    for cat in all_cats:
        code = cat.get("code") if isinstance(cat, dict) else cat[1]
        gval = cat.get("gender") if isinstance(cat, dict) else cat[2]
        key = (str(code), int(gval) if gval is not None else -1)
        if key not in present_codes:
            race_no = race_map.get(key, ("", 0))[0]
            data_rows.append([race_no, f"{code} ({gender_to_str(gval)})"]) 

    # Sort rows
    if sort_by_race:
        def key_fn(row: List[str]) -> Tuple[int, int, str]:
            race_no_cell = row[0]
            cat = str(row[1])
            m = re.match(r"^(.*?)(?:\s*\((Men|Women|Open)\))?$", cat, flags=re.IGNORECASE)
            base = (m.group(1) if m else cat).strip()
            gstr = (m.group(2).capitalize() if m and m.group(2) else 'Open')
            gval = 0 if gstr == 'Men' else (1 if gstr == 'Women' else 2)
            race_no = int(race_no_cell) if str(race_no_cell).isdigit() else 10**9
            start_off = race_map.get((base, gval), (10**9, 10**9))[1]
            return (race_no, start_off, base.lower())

        data_rows.sort(key=key_fn)
    else:
        # Alphabetical by Category (second column)
        data_rows.sort(key=lambda r: str(r[1]).lower())

    # Determine header length (max number of ranges across rows)
    max_ranges = 0
    for row in data_rows:
        max_ranges = max(max_ranges, max(0, len(row) - 1))

    wb = Workbook()
    ws = wb.active
    header = ["Race", "Category"] + [f"Range{i}" for i in range(1, max_ranges + 1)]
    ws.append(header)

    # Write rows, inserting a blank separator between races if sorting by race
    current_row = 2
    last_race_no = None
    for row in data_rows:
        race_no = row[0]
        if sort_by_race and last_race_no is not None and str(race_no) != str(last_race_no):
            current_row += 1  # blank separator row
            ws.append([None])
        ws.append(row)
        last_race_no = race_no
        current_row += 1

    # Cosmetic: set column widths, freeze top row, and enable autofilter
    max_cols = 2 + max_ranges
    # Set widths: Race=6, Category=30, Ranges=12
    for col_idx in range(1, max_cols + 1):
        col_letter = get_column_letter(col_idx)
        if col_idx == 1:
            ws.column_dimensions[col_letter].width = 6
        elif col_idx == 2:
            ws.column_dimensions[col_letter].width = 30
        else:
            ws.column_dimensions[col_letter].width = 12
    # Center align Race and Range columns
    center = Alignment(horizontal="center")
    for r in ws.iter_rows(min_row=2, min_col=1, max_col=max_cols):
        # Race column (first)
        if r[0].value is not None and str(r[0].value).strip() != "":
            r[0].alignment = center
        # Range columns start from 3rd column
        for c in range(2, max_cols):
            if r[c].value is not None and str(r[c].value).strip() != "":
                r[c].alignment = center

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(output_path)
    print(f"Wrote {ws.max_row} rows to {output_path}")


def upload_xlsx(db: RaceDBSQL, comp: Dict, input_path: str, replace: bool = False, dry_run: bool = False) -> None:
    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise SystemExit(f"openpyxl not available: {e}")

    wb = load_workbook(input_path)
    ws = wb.active

    # Build code->ranges mapping (normalized) and group by ranges
    range_to_codes: Dict[str, List[Tuple[str, Optional[str]]]] = {}
    category_col_index: Optional[int] = None
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row or all(c is None or str(c).strip() == '' for c in row):
            continue

        # Determine category column on first non-empty row
        if category_col_index is None:
            # Header cases
            lowered = [str(c).strip().lower() if c is not None else '' for c in row]
            if 'category' in lowered:
                category_col_index = lowered.index('category')
                continue  # skip header row
            # If first cell header is 'race' and second 'category'
            if len(lowered) > 1 and lowered[0] == 'race' and lowered[1] == 'category':
                category_col_index = 1
                continue  # skip header row
            # Infer: numeric first cell implies race number column present
            if str(row[0]).strip().isdigit() and len(row) > 1:
                category_col_index = 1
            else:
                category_col_index = 0

        cat_cell = row[category_col_index]
        if cat_cell is None or str(cat_cell).strip() == '':
            continue
        cat_str = str(cat_cell).strip()
        # Parse code and optional gender suffix like "(Men)"/"(Women)"/"(Open)"
        m = re.match(r"^(.*?)(?:\s*\((Men|Women|Open)\))?$", cat_str, flags=re.IGNORECASE)
        base_code = (m.group(1) if m else cat_str).strip()
        gstr = (m.group(2).capitalize() if m and m.group(2) else None)
        # Ranges start after the category column
        ranges = [str(c).strip() for c in row[category_col_index + 1:] if c is not None and str(c).strip()]
        if not ranges:
            continue
        norm = normalize_ranges(ranges)
        range_to_codes.setdefault(norm, []).append((base_code, gstr))

    if replace:
        # Delete existing links and numbers for this competition
        db.cur_execute(
            "Delete existing core_categorynumbers_categories",
            "DELETE FROM core_categorynumbers_categories WHERE categorynumbers_id IN (SELECT id FROM core_categorynumbers WHERE competition_id = %s);",
            (comp["id"],),
            debug=False,
        )
        db.cur_execute(
            "Delete existing core_categorynumbers",
            "DELETE FROM core_categorynumbers WHERE competition_id = %s;",
            (comp["id"],),
            debug=False,
        )

    # Fetch category code -> id map
    code_gender_to_id = get_code_gender_to_id(db, comp["category_format_id"])

    def gstr_to_val(g: Optional[str]) -> int:
        if not g:
            return 2
        gl = g.lower()
        if gl.startswith('men') or gl == 'm':
            return 0
        if gl.startswith('women') or gl == 'f':
            return 1
        return 2

    inserted_cn = 0
    inserted_links = 0
    deleted_links = 0
    deleted_cn = 0

    # If not replacing entirely, prune existing DB links/rows not present in desired mapping
    if not replace:
        # Build existing mapping: cn_id -> (range_str, set(category_ids))
        db.cur_execute(
            "Fetch existing categorynumbers and links",
            """
            SELECT cn.id AS cn_id, cn.range_str, cnc.category_id
            FROM core_categorynumbers cn
            LEFT JOIN core_categorynumbers_categories cnc ON cnc.categorynumbers_id = cn.id
            WHERE cn.competition_id = %s;
            """,
            (comp["id"],),
            debug=False,
        )
        rows = db.cur.fetchall() or []
        existing: Dict[int, Tuple[str, set]] = {}
        for r in rows:
            cn_id = r.get("cn_id") if isinstance(r, dict) else r[0]
            rng = r.get("range_str") if isinstance(r, dict) else r[1]
            cat_id = r.get("category_id") if isinstance(r, dict) else r[2]
            if cn_id not in existing:
                existing[cn_id] = (rng, set())
            if cat_id is not None:
                existing[cn_id][1].add(int(cat_id))

        # For each existing cn row, delete links not desired; delete empty rows with no desired categories
        for cn_id, (rng, cats) in existing.items():
            desired_pairs = []
            if rng in range_to_codes:
                for base_code, gstr in range_to_codes[rng]:
                    gval = gstr_to_val(gstr)
                    cat_id = code_gender_to_id.get((base_code, gval))
                    if cat_id:
                        desired_pairs.append(cat_id)
            desired_set = set(desired_pairs)
            # Links to remove: present in DB but not desired for this range
            for cat_id in cats - desired_set:
                db.cur_execute(
                    "Delete obsolete link",
                    "DELETE FROM core_categorynumbers_categories WHERE categorynumbers_id = %s AND category_id = %s;",
                    (cn_id, cat_id),
                    debug=False,
                )
                deleted_links += 1
            # If no desired categories for this range at all, or row became empty, delete the cn row
            # Re-check remaining links
            db.cur_execute(
                "Count links after deletion",
                "SELECT COUNT(*) FROM core_categorynumbers_categories WHERE categorynumbers_id = %s;",
                (cn_id,),
                debug=False,
            )
            cnt_row = db.cur.fetchone()
            link_count = (cnt_row[0] if not isinstance(cnt_row, dict) else list(cnt_row.values())[0]) if cnt_row is not None else 0
            if (rng not in range_to_codes) or (link_count == 0 and len(desired_set) == 0):
                # Safe to remove empty or undesired rows
                db.cur_execute(
                    "Delete empty/undesired numbers row",
                    "DELETE FROM core_categorynumbers WHERE id = %s;",
                    (cn_id,),
                    debug=False,
                )
                deleted_cn += 1
    for range_str, codes in range_to_codes.items():
        # Determine valid category IDs for this range
        valid_cat_ids: List[int] = []
        for base_code, gstr in codes:
            gval = gstr_to_val(gstr)
            cat_id = code_gender_to_id.get((base_code, gval))
            if not cat_id:
                print(f"Warning: unknown category '{base_code}' with gender '{gstr or 'Open'}' — skipping", file=sys.stderr)
                continue
            valid_cat_ids.append(cat_id)
        # If no valid categories, do not create numbers row
        if not valid_cat_ids:
            continue

        # Create or find existing numbers row for this competition and range_str
        db.cur_execute(
            "Find existing numbers row",
            "SELECT id FROM core_categorynumbers WHERE competition_id = %s AND range_str = %s;",
            (comp["id"], range_str),
            debug=False,
        )
        row = db.cur.fetchone()
        if row:
            cn_id = row.get("id") if isinstance(row, dict) else row[0]
        else:
            db.cur_execute(
                "Insert numbers row",
                "INSERT INTO core_categorynumbers (range_str, competition_id) VALUES (%s, %s) RETURNING id;",
                (range_str, comp["id"]),
                debug=False,
            )
            r = db.cur.fetchone()
            cn_id = r.get("id") if isinstance(r, dict) else r[0]
            inserted_cn += 1

        # Link to categories, avoiding duplicates
        for cat_id in valid_cat_ids:
            db.cur_execute(
                "Check existing link",
                "SELECT 1 FROM core_categorynumbers_categories WHERE categorynumbers_id = %s AND category_id = %s;",
                (cn_id, cat_id),
                debug=False,
            )
            if db.cur.fetchone():
                continue
            db.cur_execute(
                "Insert numbers link",
                "INSERT INTO core_categorynumbers_categories (categorynumbers_id, category_id) VALUES (%s, %s);",
                (cn_id, cat_id),
                debug=False,
            )
            inserted_links += 1

    if dry_run:
        db.conn.rollback()
        print(
            f"[DRY-RUN] Would insert {inserted_cn} numbers rows and {inserted_links} links"
            f"; delete {deleted_cn} numbers rows and {deleted_links} links for competition {comp['name']}"
        )
    else:
        db.conn.commit()
        print(
            f"Inserted {inserted_cn} numbers rows and {inserted_links} links"
            f"; deleted {deleted_cn} numbers rows and {deleted_links} links for competition {comp['name']}"
        )


def verify_xlsx(db: RaceDBSQL, comp: Dict, input_path: str) -> None:
    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise SystemExit(f"openpyxl not available: {e}")

    wb = load_workbook(input_path)
    ws = wb.active

    # Determine category col and build desired mapping as in upload
    range_to_codes: Dict[str, List[Tuple[str, Optional[str]]]] = {}
    category_col_index: Optional[int] = None
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row or all(c is None or str(c).strip() == '' for c in row):
            continue
        if category_col_index is None:
            lowered = [str(c).strip().lower() if c is not None else '' for c in row]
            if 'category' in lowered:
                category_col_index = lowered.index('category')
                continue
            if len(lowered) > 1 and lowered[0] == 'race' and lowered[1] == 'category':
                category_col_index = 1
                continue
            if str(row[0]).strip().isdigit() and len(row) > 1:
                category_col_index = 1
            else:
                category_col_index = 0
        cat_cell = row[category_col_index]
        if cat_cell is None or str(cat_cell).strip() == '':
            continue
        cat_str = str(cat_cell).strip()
        m = re.match(r"^(.*?)(?:\s*\((Men|Women|Open)\))?$", cat_str, flags=re.IGNORECASE)
        base_code = (m.group(1) if m else cat_str).strip()
        gstr = (m.group(2).capitalize() if m and m.group(2) else None)
        ranges = [str(c).strip() for c in row[category_col_index + 1:] if c is not None and str(c).strip()]
        if not ranges:
            continue
        norm = normalize_ranges(ranges)
        range_to_codes.setdefault(norm, []).append((base_code, gstr))

    # Build existing mapping from DB
    db.cur_execute(
        "Fetch existing categorynumbers and links",
        """
        SELECT cn.id AS cn_id, cn.range_str, cnc.category_id
        FROM core_categorynumbers cn
        LEFT JOIN core_categorynumbers_categories cnc ON cnc.categorynumbers_id = cn.id
        WHERE cn.competition_id = %s;
        """,
        (comp["id"],),
        debug=False,
    )
    rows = db.cur.fetchall() or []
    existing: Dict[str, set] = {}
    for r in rows:
        rng = r.get("range_str") if isinstance(r, dict) else r[1]
        cat_id = r.get("category_id") if isinstance(r, dict) else r[2]
        existing.setdefault(rng, set())
        if cat_id is not None:
            existing[rng].add(int(cat_id))

    code_gender_to_id = get_code_gender_to_id(db, comp["category_format_id"])
    id_to_code_gender = get_id_to_code_gender(db, comp["category_format_id"])

    def gstr_to_val(g: Optional[str]) -> int:
        if not g:
            return 2
        gl = g.lower()
        if gl.startswith('men') or gl == 'm':
            return 0
        if gl.startswith('women') or gl == 'f':
            return 1
        return 2

    # Compute desired id-sets
    desired: Dict[str, set] = {}
    for rng, pairs in range_to_codes.items():
        for base_code, gstr in pairs:
            gval = gstr_to_val(gstr)
            cat_id = code_gender_to_id.get((base_code, gval))
            if cat_id:
                desired.setdefault(rng, set()).add(cat_id)

    # Compute diffs
    inserts_rows = 0
    inserts_links: List[Tuple[str, str, str]] = []  # (range, code, gender)
    deletes_rows = 0
    deletes_links: List[Tuple[str, str, str]] = []

    # Links to add
    for rng, dset in desired.items():
        eset = existing.get(rng, set())
        for cid in dset - eset:
            code, gval = id_to_code_gender.get(cid, ("?", 2))
            gname = 'Men' if gval == 0 else ('Women' if gval == 1 else 'Open')
            inserts_links.append((rng, code, gname))
        if rng not in existing:
            inserts_rows += 1

    # Links/rows to delete
    for rng, eset in existing.items():
        dset = desired.get(rng, set())
        for cid in eset - dset:
            code, gval = id_to_code_gender.get(cid, ("?", 2))
            gname = 'Men' if gval == 0 else ('Women' if gval == 1 else 'Open')
            deletes_links.append((rng, code, gname))
        if rng not in desired:
            deletes_rows += 1

    # Print summary
    print(f"Verify: would insert {inserts_rows} numbers rows and {len(inserts_links)} links; delete {deletes_rows} numbers rows and {len(deletes_links)} links")
    if inserts_links:
        print("  Inserts:")
        for rng, code, g in inserts_links:
            print(f"    + {code} ({g}) -> {rng}")
    if deletes_links:
        print("  Deletes:")
        for rng, code, g in deletes_links:
            print(f"    - {code} ({g}) -> {rng}")

    # Close-range warnings within the same race (based on uploaded desired mapping)
    race_map = fetch_race_category_map(db, comp["id"])  # (code,g) -> (race_no, start_off)

    def last_two_digits_covered(rng_str: str) -> set:
        covered = set()
        for part in (rng_str or '').split(','):
            part = part.strip()
            if not part:
                continue
            try:
                a, b = part.split('-')
                a = int(a)
                b = int(b)
                if a > b:
                    a, b = b, a
                for x in range(a, b + 1):
                    covered.add(x % 100)
            except Exception:
                continue
        return covered

    # Build desired ranges per race
    print("  Close-range warnings:")
    race_to_ranges: Dict[int, List[Tuple[str, set]]] = {}
    for rng, pairs in range_to_codes.items():
        cov = last_two_digits_covered(rng)
        for base_code, gstr in pairs:
            gval = gstr_to_val(gstr)
            race_no = race_map.get((base_code, gval), (None, None))[0]
            if race_no is None:
                continue
            race_to_ranges.setdefault(race_no, []).append((f"{base_code} ({'Men' if gval==0 else 'Women' if gval==1 else 'Open'})", cov))
    for race_no, items in sorted(race_to_ranges.items()):
        n = len(items)
        for i in range(n):
            name_i, cov_i = items[i]
            for j in range(i + 1, n):
                name_j, cov_j = items[j]
                if cov_i and cov_j and cov_i.intersection(cov_j):
                    print(f"    Race {race_no}: potential overlap {name_i} <-> {name_j}")


def main():
    ap = argparse.ArgumentParser(description="Download or upload competition category number ranges to/from XLSX")
    sub = ap.add_subparsers(dest="cmd", required=True)

    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--host", default="localhost", help="Database host")
    common.add_argument("--name", default=None, help="Competition name")
    common.add_argument("--date", default=None, help="Competition date (YYYY-MM-DD)")

    dl = sub.add_parser("download", parents=[common], help="Download category numbers to XLSX")
    dl.add_argument("--xlsx", required=False, help="XLSX file path (defaults to sanitized competition name)")
    dl.add_argument("--sort_by_race", action="store_true", help="Sort by event start, start offset, then category code")
    dl.add_argument("--debug", action="store_true", help="Print race/start/category details during download")

    ul = sub.add_parser("upload", parents=[common], help="Upload category numbers from XLSX")
    ul.add_argument("--xlsx", required=False, help="XLSX file path (defaults to sanitized competition name)")
    ul.add_argument("--replace", action="store_true", help="Replace all existing numbers for the competition")
    ul.add_argument("--dry-run", action="store_true", help="Validate and show changes without writing")
    # Accept but ignore these options on upload for convenience (history reuse)
    ul.add_argument("--sort_by_race", action="store_true", help=argparse.SUPPRESS)
    ul.add_argument("--debug", action="store_true", help=argparse.SUPPRESS)

    vr = sub.add_parser("verify", parents=[common], help="Show differences between DB and XLSX without writing; warns about close-range overlaps")
    vr.add_argument("--xlsx", required=False, help="XLSX file path (defaults to sanitized competition name)")

    args = ap.parse_args()

    db = RaceDBSQL(host=args.host)
    comp = find_competition(db, args.name, args.date)

    if args.cmd == "download":
        xlsx = args.xlsx or f"{sanitize_filename(comp.get('name','competition'))}.xlsx"
        download_xlsx(db, comp, xlsx, sort_by_race=args.sort_by_race, debug=getattr(args, 'debug', False))
    elif args.cmd == "upload":
        xlsx = args.xlsx or f"{sanitize_filename(comp.get('name','competition'))}.xlsx"
        upload_xlsx(db, comp, xlsx, replace=args.replace, dry_run=args.dry_run)
    elif args.cmd == "verify":
        xlsx = getattr(args, 'xlsx', None) or f"{sanitize_filename(comp.get('name','competition'))}.xlsx"
        verify_xlsx(db, comp, xlsx)


if __name__ == "__main__":
    main()
